"""
Unit tests for PAIP DocxFormatInspector and DocxAutoFormatter.
"""

from io import BytesIO
from docx import Document
from core.document.format_inspector import docx_format_inspector, docx_auto_formatter


def test_docx_format_inspection_detects_non_standard_fonts_and_tables():
    """Kiểm tra khả năng bắt lỗi font trong cả đoạn văn và bảng biểu."""
    doc = Document()
    
    # Đoạn văn ngoài bảng bị sai font
    p1 = doc.add_paragraph()
    r1 = p1.add_run("Văn bản dùng font Arial ngoài thân bài.")
    r1.font.name = "Arial"

    # Bảng biểu bị sai font
    table = doc.add_table(rows=2, cols=2)
    cell_p = table.rows[0].cells[0].paragraphs[0]
    r_cell = cell_p.add_run("Văn bản dùng font Calibri trong bảng.")
    r_cell.font.name = "Calibri"

    # Đoạn văn dài chưa căn đều
    doc.add_paragraph("Đoạn văn này được viết rất dài nhằm mục đích kiểm tra khả năng phát hiện lỗi căn lề trái thay vì căn đều hai bên theo Nghị định 30.")

    buf = BytesIO()
    doc.save(buf)
    raw_bytes = buf.getvalue()

    report = docx_format_inspector.inspect(raw_bytes)
    assert report["success"] is True
    assert report["font_clean"] is False
    assert len(report["font_issues"]) >= 2
    assert report["alignment_clean"] is False
    assert len(report["alignment_issues"]) >= 1


def test_docx_auto_formatter_fixes_all_issues():
    """Kiểm tra auto-formatter sửa triệt để font và căn lề."""
    doc = Document()
    
    p1 = doc.add_paragraph()
    r1 = p1.add_run("Đoạn văn có font Tahoma.")
    r1.font.name = "Tahoma"

    table = doc.add_table(rows=1, cols=1)
    r_cell = table.rows[0].cells[0].paragraphs[0].add_run("Ô bảng có font Segoe UI.")
    r_cell.font.name = "Segoe UI"

    buf = BytesIO()
    doc.save(buf)
    raw_bytes = buf.getvalue()

    # Chuẩn hóa tự động
    fixed_stream = docx_auto_formatter.auto_format(raw_bytes)
    after_report = docx_format_inspector.inspect(fixed_stream.getvalue())

    assert after_report["font_clean"] is True
    assert len(after_report["font_issues"]) == 0
    assert after_report["margins_compliant"] is True
    assert after_report["margins"]["left_cm"] == 3.0
    assert after_report["margins"]["right_cm"] == 1.5
    assert after_report["margins"]["top_cm"] == 2.0
    assert after_report["margins"]["bottom_cm"] == 2.0


def test_docx_theme_font_detection_and_repair():
    """Kiểm tra phát hiện và chuẩn hóa font ẩn dưới dạng XML Theme (w:asciiTheme=minorHAnsi)."""
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("Đoạn văn dùng theme font Calibri.")
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:asciiTheme'), 'minorHAnsi')
    rFonts.set(qn('w:hAnsiTheme'), 'minorHAnsi')

    buf = BytesIO()
    doc.save(buf)
    raw_bytes = buf.getvalue()

    report = docx_format_inspector.inspect(raw_bytes)
    assert report["font_clean"] is False
    assert len(report["font_issues"]) >= 1
    assert report["font_issues"][0]["font"].lower() != "times new roman"

    # Auto fix
    fixed_stream = docx_auto_formatter.auto_format(raw_bytes)
    after_report = docx_format_inspector.inspect(fixed_stream.getvalue())
    assert after_report["font_clean"] is True
    assert len(after_report["font_issues"]) == 0


