"""
Tests for Agent 0 — Document Proofreader & PAIP Core.

Test cases thiết kế dựa trên POC note trong Obsidian:
    PTSC_AI_RnD/Lab&Research/POC/POC_001_Document_Proofreader.md
"""

import pytest
from httpx import ASGITransport, AsyncClient

from api.main import app
from core.common.schemas import LLMProvider
from core.document.reader import document_reader


@pytest.fixture
def client():
    """Create async test client."""
    transport = ASGITransport(app=app)
    return AsyncClient(transport=transport, base_url="http://test")


# ── Unit Tests: Schemas ───────────────────────────────────

class TestProofreadSchemas:
    """Test schemas validation."""

    def test_proofread_response_success(self):
        from agents.agent_0_proofreader.schemas import ProofreadResponse

        resp = ProofreadResponse(success=True, message="OK")
        assert resp.success is True
        assert resp.processing_time_ms == 0.0

    def test_proofread_response_with_errors(self):
        from agents.agent_0_proofreader.schemas import (
            ErrorType,
            ProofreadError,
            ProofreadResult,
            Severity,
        )

        error = ProofreadError(
            type=ErrorType.SPELLING,
            original="tôi ddang",
            suggested="tôi đang",
            explanation="Lỗi đánh máy: thừa ký tự 'd'",
            severity=Severity.MEDIUM,
        )
        result = ProofreadResult(total_errors=1, errors=[error], score=8.5)
        assert result.total_errors == 1
        assert len(result.errors) == 1
        assert result.errors[0].type == ErrorType.SPELLING
        assert result.score == 8.5


# ── Unit Tests: Document Reader ───────────────────────────

class TestDocumentReader:
    """Test DocumentReader behavior."""

    def test_read_non_existent_file_raises(self):
        with pytest.raises(FileNotFoundError):
            document_reader.read("non_existent_file_12345.docx")

    def test_read_unsupported_extension_raises(self, tmp_path):
        bad_file = tmp_path / "test.xyz"
        bad_file.write_text("hello", encoding="utf-8")
        with pytest.raises(ValueError) as exc_info:
            document_reader.read(bad_file)
        assert "không được hỗ trợ" in str(exc_info.value)

    def test_read_text_and_markdown(self, tmp_path):
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("Xin chào PTSC", encoding="utf-8")
        assert document_reader.read(txt_file) == "Xin chào PTSC"

        md_file = tmp_path / "test.md"
        md_file.write_text("# Tiêu đề", encoding="utf-8")
        assert document_reader.read(md_file) == "# Tiêu đề"


# ── Unit Tests: Service Logic & Parsing ───────────────────

class TestProofreadService:
    """Test service logic and parsing resilience."""

    @pytest.mark.asyncio
    async def test_empty_text_returns_error(self):
        from agents.agent_0_proofreader.service import proofreader_service

        result = await proofreader_service.proofread_text("")
        assert result.success is False
        assert "trống" in result.message

    @pytest.mark.asyncio
    async def test_text_too_long_returns_error(self):
        from agents.agent_0_proofreader.service import proofreader_service

        long_text = "a" * 60_000
        result = await proofreader_service.proofread_text(long_text)
        assert result.success is False
        assert "quá dài" in result.message

    def test_extract_json_with_fenced_code_block(self):
        from agents.agent_0_proofreader.service import proofreader_service

        llm_raw = """
        Dưới đây là kết quả phân tích:
        ```json
        {
          "total_errors": 1,
          "errors": [
            {
              "type": "spelling",
              "original": "ngayf",
              "suggested": "ngày",
              "explanation": "Lỗi gõ Telex",
              "severity": "medium"
            }
          ],
          "summary": "Tốt",
          "score": 9.0
        }
        ```
        Hy vọng giúp ích cho bạn!
        """
        result = proofreader_service._parse_llm_response(llm_raw)
        assert result.total_errors == 1
        assert len(result.errors) == 1
        assert result.errors[0].original == "ngayf"
        assert result.score == 9.0

    def test_sanitize_error_with_variations(self):
        from agents.agent_0_proofreader.schemas import ErrorType, Severity
        from agents.agent_0_proofreader.service import proofreader_service

        err1 = proofreader_service._sanitize_error({"type": "chính tả", "severity": "cao"})
        assert err1.type == ErrorType.SPELLING
        assert err1.severity == Severity.HIGH

        err2 = proofreader_service._sanitize_error({"type": "grammar_error", "severity": "thấp"})
        assert err2.type == ErrorType.GRAMMAR
        assert err2.severity == Severity.LOW


# ── API Tests ─────────────────────────────────────────────

class TestProofreadAPI:
    """Test API endpoints."""

    @pytest.mark.asyncio
    async def test_health_check(self, client):
        async with client:
            resp_root = await client.get("/")
            assert resp_root.status_code == 200

            resp_health = await client.get("/health")
            assert resp_health.status_code == 200
            data = resp_health.json()
            assert data["status"] == "ok"
            assert data["app_name"] == "PAIP"

    @pytest.mark.asyncio
    async def test_proofread_empty_text_json(self, client):
        async with client:
            resp = await client.post(
                "/api/v1/proofread/text",
                json={"document_text": ""},
            )
            assert resp.status_code == 200
            data = resp.json()
            assert data["success"] is False
            assert "trống" in data["message"]

    @pytest.mark.asyncio
    async def test_export_docx_endpoint(self, client):
        async with client:
            resp = await client.post(
                "/api/v1/proofread/export-docx",
                json={
                    "text": "# HỢP ĐỒNG KINH TẾ\n\nCỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n\nĐiều 1: Nội dung công việc.\n- Cung cấp 08 laptop Dell.",
                    "filename": "test_contract.docx",
                    "title": "Hợp Đồng Mua Sắm",
                },
            )
            assert resp.status_code == 200
            assert "wordprocessingml.document" in resp.headers.get("content-type", "")
            assert len(resp.content) > 1000

    @pytest.mark.asyncio
    async def test_export_docx_inplace_endpoint(self, client):
        from docx import Document
        from io import BytesIO
        import json

        # Create a sample docx in memory
        doc = Document()
        p = doc.add_paragraph("Kính gởi Ban giám đốc cần bổ xung nhân sự.")
        buf = BytesIO()
        doc.save(buf)
        raw_docx = buf.getvalue()

        async with client:
            resp = await client.post(
                "/api/v1/proofread/export-docx-inplace",
                files={"file": ("test_origin.docx", raw_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={
                    "replacements": json.dumps([
                        {"original": "Kính gởi", "suggested": "Kính gửi"},
                        {"original": "bổ xung", "suggested": "bổ sung"},
                    ]),
                    "highlight_changes": "false",
                },
            )
            assert resp.status_code == 200
            assert "wordprocessingml.document" in resp.headers.get("content-type", "")
            
            # Verify the downloaded docx actually has the replaced text
            doc_res = Document(BytesIO(resp.content))
            assert "Kính gửi Ban giám đốc cần bổ sung nhân sự." in doc_res.paragraphs[0].text


# ── Unit Tests: Document Writer Formatting Preservation ──

class TestDocumentWriterFormatPreservation:
    """Test bảo toàn tuyệt đối 100% định dạng và font chữ khi sửa lỗi in-place."""

    def test_multi_run_format_preservation(self):
        from docx import Document
        from core.document.writer import DocumentWriter

        doc = Document()
        p = doc.add_paragraph()

        # Tạo đoạn văn có 3 run với 3 định dạng khác nhau:
        # Run 0: "Kính " (thường)
        # Run 1: "gởi " (in nghiêng)
        # Run 2: "Ban Giám Đốc" (in đậm, font Arial)
        r0 = p.add_run("Kính ")
        r1 = p.add_run("gởi ")
        r1.italic = True
        r2 = p.add_run("Ban Giám Đốc")
        r2.bold = True
        r2.font.name = "Arial"

        # Thay thế "Kính gởi" -> "Kính gửi" (chuỗi trải dài qua run 0 và run 1)
        replaced = DocumentWriter._replace_text_in_paragraph(p, "Kính gởi", "Kính gửi")
        assert replaced is True

        # Kiểm tra nội dung toàn đoạn
        assert p.text == "Kính gửi Ban Giám Đốc"

        # Kiểm tra tính bảo toàn định dạng của Run 2 (Ban Giám Đốc)
        assert r2.text == "Ban Giám Đốc"
        assert r2.bold is True
        assert r2.font.name == "Arial"

    def test_replace_in_table_and_header_preservation(self):
        from docx import Document
        from io import BytesIO
        from core.document.writer import document_writer

        doc = Document()
        
        # Thêm table
        table = doc.add_table(rows=2, cols=2)
        cell = table.cell(0, 1)
        p_cell = cell.paragraphs[0]
        r_norm = p_cell.add_run("Dự án ")
        r_bold = p_cell.add_run("triễn khai ")
        r_bold.bold = True
        r_ital = p_cell.add_run("khẩn cấp.")
        r_ital.italic = True

        buf = BytesIO()
        doc.save(buf)
        raw_docx = buf.getvalue()

        # Áp dụng in-place
        out_stream = document_writer.replace_in_docx(
            doc_bytes=raw_docx,
            replacements=[{"original": "triễn khai", "suggested": "triển khai"}],
        )

        res_doc = Document(out_stream)
        cell_res = res_doc.tables[0].cell(0, 1)
        assert "Dự án triển khai khẩn cấp." in cell_res.text
        # Kiểm tra run in nghiêng cuối cùng vẫn giữ nguyên
        last_run = cell_res.paragraphs[0].runs[-1]
        assert last_run.italic is True
        assert "khẩn cấp." in last_run.text


# ── Unit Tests: Legal Verification & Real-time Grounding ──

class TestLegalSearchGrounding:
    """Test tính năng thẩm định pháp lý và tra cứu văn bản quy phạm pháp luật."""

    @pytest.mark.asyncio
    async def test_legal_mode_detects_law_and_returns_legal_type(self):
        from agents.agent_0_proofreader.service import proofreader_service
        from agents.agent_0_proofreader.schemas import ErrorType

        doc_text = "Căn cứ theo Luật Chuyển đổi số và Luật Đấu thầu 2013 để tiến hành lựa chọn nhà thầu."
        res = await proofreader_service.proofread_text(
            text=doc_text,
            mode="legal",
            provider=LLMProvider.MOCK,
        )

        assert res.success is True
        assert res.result is not None
        assert res.result.total_errors > 0

        # Kiểm tra có ít nhất 1 lỗi thuộc ErrorType.LEGAL
        legal_errors = [e for e in res.result.errors if e.type == ErrorType.LEGAL]
        assert len(legal_errors) >= 1

        # Kiểm tra nội dung trích dẫn & căn cứ pháp lý
        first_legal = legal_errors[0]
        assert first_legal.reference is not None or "Quy định" in first_legal.explanation or "Luật" in first_legal.explanation

    @pytest.mark.asyncio
    async def test_custom_instructions_law_keyword_triggers_search(self):
        from agents.agent_0_proofreader.service import proofreader_service
        from agents.agent_0_proofreader.schemas import ErrorType

        doc_text = "Thực hiện theo Luật Đấu thầu số 43/2013/QH13."
        res = await proofreader_service.proofread_text(
            text=doc_text,
            mode="standard",
            custom_instructions="Kiểm tra xem các tên riêng của các Luật trong này đã đúng chưa, các quyết định đã đúng với quy định pháp luật hiện hành chưa",
            provider=LLMProvider.MOCK,
        )

        assert res.success is True
        assert res.result is not None
        assert any(e.type == ErrorType.LEGAL for e in res.result.errors)

    def test_sanitize_legal_error_preserves_reference_and_source(self):
        from agents.agent_0_proofreader.service import ProofreaderService
        from agents.agent_0_proofreader.schemas import ErrorType

        raw_item = {
            "type": "pháp lý",
            "original": "Luật Đấu thầu 2013",
            "suggested": "Luật Đấu thầu 2023",
            "explanation": "Đã hết hiệu lực từ 01/01/2024",
            "severity": "high",
            "reference": "Luật số 22/2023/QH15",
            "source_link": "https://thuvienphapluat.vn",
        }

        err = ProofreaderService._sanitize_error(raw_item)
        assert err.type == ErrorType.LEGAL
        assert err.reference == "Luật số 22/2023/QH15"
        assert err.source_link == "https://thuvienphapluat.vn"
