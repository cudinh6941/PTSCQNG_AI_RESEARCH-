"""
Unit & Integration Tests for DocxFormatInspector and DocxAutoFormatter.
"""

from io import BytesIO
import pytest
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from core.document.format_inspector import DocxFormatInspector, DocxAutoFormatter


def create_sample_docx_with_bad_format() -> bytes:
    """Tạo file docx giả lập có lề sai và font không đồng nhất."""
    doc = Document()
    
    # Set non-standard margins (e.g. 1.0 cm)
    for section in doc.sections:
        section.top_margin = Cm(1.0)
        section.bottom_margin = Cm(1.0)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.0)

    # Add paragraphs with non-standard font & alignment
    p1 = doc.add_paragraph()
    r1 = p1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    r1.font.name = "Arial"

    p2 = doc.add_paragraph()
    r2 = p2.add_run(
        "Kính gửi Ban Giám đốc Tổng công ty Cổ phần Dịch vụ Kỹ thuật Dầu khí Việt Nam. "
        "Chúng tôi xin trân trọng báo cáo toàn bộ tiến độ triển khai các hạng mục thi công "
        "chế tạo chân đế giàn khoan ngoài khơi thuộc dự án Lô B Ô Môn theo đúng kế hoạch."
    )
    r2.font.name = "Calibri"
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT

    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_format_inspector_detects_bad_margins():
    bad_doc_bytes = create_sample_docx_with_bad_format()
    report = DocxFormatInspector.inspect(bad_doc_bytes)

    assert report["success"] is True
    assert report["margins_compliant"] is False
    assert report["margins"]["top_cm"] == 1.0
    assert report["margins"]["left_cm"] == 1.5


def test_docx_auto_formatter():
    bad_doc_bytes = create_sample_docx_with_bad_format()
    fixed_stream = DocxAutoFormatter.auto_format(bad_doc_bytes)
    fixed_bytes = fixed_stream.getvalue()

    # Re-inspect fixed document
    report = DocxFormatInspector.inspect(fixed_bytes)
    assert report["success"] is True
    assert report["margins_compliant"] is True
    assert report["margins"]["top_cm"] == 2.0
    assert report["margins"]["bottom_cm"] == 2.0
    assert report["margins"]["left_cm"] == 3.0
    assert report["margins"]["right_cm"] == 1.5
    assert report["dominant_font"] == "Times New Roman"
    assert report["alignment_clean"] is True
